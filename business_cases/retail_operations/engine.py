"""Two usable execution routes for an exact signed retail ledger report."""
import argparse
import json
import sqlite3
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent


def aggregate(rows, period):
    raw = excluded = eligible = gross = credits = cancel = other = 0
    invoices, customers = set(), defaultdict(set)
    countries = defaultdict(lambda: [0, 0, 0])
    for invoice, quantity, month, price, customer, country in rows:
        if month != period:
            continue
        raw += 1
        if price <= 0 or quantity == 0:
            excluded += 1
            continue
        eligible += 1
        amount = quantity * price
        countries[country][0] += amount
        if amount > 0:
            gross += amount
            countries[country][1] += amount
            if not invoice.upper().startswith('C'):
                invoices.add(invoice)
                if customer:
                    customers[customer].add(invoice)
        else:
            credits -= amount
            countries[country][2] -= amount
            if invoice.upper().startswith('C'):
                cancel -= amount
            else:
                other -= amount
    return {'period': period, 'raw_rows': raw, 'excluded_rows': excluded,
            'eligible_rows': eligible, 'gross_positive_micros': gross,
            'credit_abs_micros': credits, 'net_micros': gross - credits,
            'cancel_credit_micros': cancel, 'noncancel_negative_micros': other,
            'positive_order_count': len(invoices), 'known_customer_count': len(customers),
            'repeat_customer_count': sum(len(v) >= 2 for v in customers.values()),
            'countries': {k: {'net_micros': v[0], 'gross_positive_micros': v[1],
                              'credit_abs_micros': v[2]} for k, v in sorted(countries.items())}}


def raw_rows():
    from openpyxl import load_workbook
    book = load_workbook(ROOT / 'data/raw/online_retail.xlsx', read_only=True, data_only=True)
    try:
        for invoice, stock, description, quantity, date, price, customer, country in book.active.iter_rows(min_row=2, values_only=True):
            yield (str(invoice), int(quantity), date.strftime('%Y-%m'),
                   int((Decimal(str(price)) * 1000000).to_integral_value()), str(int(customer)) if customer is not None else '',
                   str(country))
    finally:
        book.close()


def sqlite_report(period):
    with sqlite3.connect(f'file:{ROOT / "data/ledger.sqlite3"}?mode=ro', uri=True) as db:
        rows = db.execute('SELECT invoice,quantity,month,price_micros,customer,country FROM ledger WHERE month=?', (period,))
        return aggregate(rows, period)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--engine', choices=['raw_xlsx', 'indexed_sqlite'], required=True)
    parser.add_argument('--period', required=True)
    args = parser.parse_args()
    result = aggregate(raw_rows(), args.period) if args.engine == 'raw_xlsx' else sqlite_report(args.period)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
